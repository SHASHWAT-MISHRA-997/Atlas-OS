"""Document search: full-text index of files in the agent's approved folders.

Indexing reads file contents, so it runs only as an approved task (kind ``index_documents``,
see app/agent.py). It is local: text goes into ATLAS's SQLite database with an FTS5 index,
and no embedding model is downloaded. Search ranks by BM25.

Text from documents is shown to the model only when the user turns on "Use my documents" in
chat, and it is marked there as untrusted reference text, never as instructions.
"""
from __future__ import annotations

import os
import re
import sqlite3
import time
from pathlib import Path

from . import storage

TEXT_TYPES = {".txt", ".md", ".markdown", ".csv", ".log", ".json"}
SUPPORTED = TEXT_TYPES | {".pdf", ".docx"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILES = 5000
MAX_DEPTH = 8
MAX_CHUNKS_PER_FILE = 400
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
INDEX_SECONDS = 300.0
STOPWORDS = {"the", "a", "an", "and", "or", "of", "to", "in", "is", "are", "was", "for", "on", "with", "what",
             "which", "who", "how", "my", "me", "i", "it", "this", "that", "do", "does", "kya", "hai", "ka", "ki",
             "ke", "mera", "meri", "mere", "se", "ko", "aur"}


class UnreadableFile(Exception):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_TYPES:
            data = path.read_bytes()
            for encoding in ("utf-8-sig", "utf-16", "cp1252"):
                try:
                    return data.decode(encoding)
                except UnicodeDecodeError:
                    continue
            return data.decode("utf-8", errors="replace")
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise UnreadableFile("password-protected PDF")
            return "\n".join((page.extract_text() or "") for page in reader.pages[:500])
        if suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
    except UnreadableFile:
        raise
    except Exception as exc:  # parsers raise many types on damaged files
        raise UnreadableFile(exc.__class__.__name__) from None
    raise UnreadableFile("unsupported type")


def chunks_of(text: str) -> list[str]:
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    out, start = [], 0
    while start < len(text) and len(out) < MAX_CHUNKS_PER_FILE:
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):
            # Prefer to break at a paragraph or sentence boundary.
            cut = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
            if cut > start + CHUNK_CHARS // 2:
                end = cut + 1
        piece = text[start:end].strip()
        if piece:
            out.append(piece)
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return out


def _walk(root: Path):
    """Supported files under root, never following links, with limits."""
    base = len(root.parts)
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False, onerror=lambda _e: None):
        here = Path(dirpath)
        dirnames[:] = [d for d in dirnames if not d.startswith(".")] if len(here.parts) - base < MAX_DEPTH else []
        for name in filenames:
            path = here / name
            if path.suffix.lower() in SUPPORTED and not name.startswith(("~$", ".")) and not path.is_symlink():
                yield path


def index_folder(folder: str, root: Path) -> dict:
    """Bring the index for one approved folder up to date. Unchanged files are skipped."""
    deadline = time.monotonic() + INDEX_SECONDS
    seen: set[str] = set()
    counts = {"indexed": 0, "unchanged": 0, "skipped": 0, "failed": 0, "removed": 0}
    stopped = None
    with storage.connection() as conn:
        known = {r["path"]: (r["size"], r["mtime"]) for r in
                 conn.execute("SELECT path, size, mtime FROM doc_files WHERE folder = ?", (folder,))}
        for path in _walk(root):
            if len(seen) >= MAX_FILES:
                stopped = "file limit"
                break
            if time.monotonic() > deadline:
                stopped = "time limit"
                break
            rel = path.relative_to(root).as_posix()
            seen.add(rel)
            try:
                st = path.stat()
            except OSError:
                counts["failed"] += 1
                continue
            if st.st_size > MAX_FILE_BYTES:
                counts["skipped"] += 1
                continue
            if known.get(rel) == (st.st_size, st.st_mtime):
                counts["unchanged"] += 1
                continue
            error, pieces = None, []
            try:
                pieces = chunks_of(extract_text(path))
            except UnreadableFile as exc:
                error = str(exc)
            _store(conn, folder, rel, st.st_size, st.st_mtime, pieces, error)
            counts["failed" if error else "indexed"] += 1
        if stopped is None:
            # Only a complete walk can tell which files were deleted.
            for rel in set(known) - seen:
                _remove(conn, folder, rel)
                counts["removed"] += 1
    return {"folder": folder, **counts, "stopped_at": stopped, **status_of(folder)}


def _store(conn: sqlite3.Connection, folder: str, rel: str, size: int, mtime: float, pieces: list[str], error: str | None):
    with conn:
        conn.execute("DELETE FROM doc_chunks WHERE folder = ? AND path = ?", (folder, rel))
        conn.executemany("INSERT INTO doc_chunks(folder, path, chunk, text) VALUES (?, ?, ?, ?)",
                         [(folder, rel, i, piece) for i, piece in enumerate(pieces)])
        conn.execute("INSERT INTO doc_files(folder, path, size, mtime, chunks, error, indexed_at) VALUES (?, ?, ?, ?, ?, ?, ?) "
                     "ON CONFLICT(folder, path) DO UPDATE SET size = excluded.size, mtime = excluded.mtime, "
                     "chunks = excluded.chunks, error = excluded.error, indexed_at = excluded.indexed_at",
                     (folder, rel, size, mtime, len(pieces), error, storage._now()))


def _remove(conn: sqlite3.Connection, folder: str, rel: str) -> None:
    with conn:
        conn.execute("DELETE FROM doc_chunks WHERE folder = ? AND path = ?", (folder, rel))
        conn.execute("DELETE FROM doc_files WHERE folder = ? AND path = ?", (folder, rel))


def status_of(folder: str | None = None) -> dict:
    with storage.connection() as conn:
        where, args = ("WHERE folder = ?", (folder,)) if folder else ("", ())
        row = conn.execute(f"SELECT COUNT(*) AS files, COALESCE(SUM(chunks), 0) AS chunks, "  # noqa: S608 - fixed clause
                           f"SUM(error IS NOT NULL) AS unreadable, MAX(indexed_at) AS last FROM doc_files {where}", args).fetchone()
    return {"files": row["files"], "chunks": row["chunks"], "unreadable": row["unreadable"] or 0, "last_indexed": row["last"]}


def status() -> dict:
    with storage.connection() as conn:
        folders = [r["folder"] for r in conn.execute("SELECT DISTINCT folder FROM doc_files ORDER BY folder")]
    return {"folders": {f: status_of(f) for f in folders}, **status_of()}


def _query(q: str) -> str | None:
    words = [w for w in re.findall(r"\w+", q.lower(), flags=re.UNICODE) if len(w) > 1 and w not in STOPWORDS][:12]
    return " OR ".join(f'"{w}"*' for w in words) if words else None


def search(q: str, limit: int = 5, folder: str | None = None) -> list[dict]:
    query = _query(q)
    if not query:
        return []
    sql = ("SELECT c.folder, c.path, c.chunk, c.text, bm25(doc_chunks_fts) AS score, "
           "snippet(doc_chunks_fts, 0, '[', ']', '…', 20) AS snippet "
           "FROM doc_chunks_fts JOIN doc_chunks c ON c.id = doc_chunks_fts.rowid WHERE doc_chunks_fts MATCH ?")
    args: list = [query]
    if folder:
        sql += " AND c.folder = ?"
        args.append(folder)
    sql += " ORDER BY score LIMIT ?"
    args.append(max(1, min(limit, 20)))
    with storage.connection() as conn:
        rows = conn.execute(sql, args).fetchall()
    return [{"folder": r["folder"], "path": r["path"], "chunk": r["chunk"], "text": r["text"],
             "snippet": r["snippet"], "score": round(-r["score"], 3)} for r in rows]


def clear(folder: str | None = None) -> int:
    with storage.connection() as conn, conn:
        where, args = ("WHERE folder = ?", (folder,)) if folder else ("", ())
        conn.execute(f"DELETE FROM doc_chunks {where}", args)  # noqa: S608 - fixed clause
        return conn.execute(f"DELETE FROM doc_files {where}", args).rowcount  # noqa: S608 - fixed clause
