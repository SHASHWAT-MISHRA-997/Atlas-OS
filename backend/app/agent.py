"""Desktop agent: a fixed set of typed actions on the user's own PC.

There is no shell and no free-form command. Each action has a schema, a risk
note shown to the user, and runs only as a background task that the user has
approved (see app/task_engine.py). Everything is off unless
ATLAS_ENABLE_DESKTOP_AGENT=true, which is meant for the backend running on the
user's own computer.

File actions are limited to approved folders. By default those are Documents,
Downloads and Desktop in the user's home folder; set
ATLAS_AGENT_FOLDERS="name=path;name=path" to choose others. Paths are resolved
(symlinks included) and must stay inside their folder.

Writes are deliberately narrow (D-015): new files only, never overwriting, and
"delete" moves a file into ATLAS's own trash folder, from which it can be
restored. Nothing is ever deleted permanently by an action.
"""
from __future__ import annotations

import json
import os
import platform
import re
import shutil
import subprocess
import time
import uuid
from datetime import UTC, datetime
from pathlib import Path

from . import docsearch, storage
from .config import get_settings

APPS = {
    "notepad": {"label": "Notepad", "commands": {"Windows": ["notepad.exe"], "Linux": ["gedit"], "Darwin": ["open", "-a", "TextEdit"]}},
    "calculator": {"label": "Calculator", "commands": {"Windows": ["calc.exe"], "Linux": ["gnome-calculator"], "Darwin": ["open", "-a", "Calculator"]}},
    "file_manager": {"label": "File Manager", "commands": {"Windows": ["explorer.exe"], "Linux": ["xdg-open", str(Path.home())], "Darwin": ["open", str(Path.home())]}},
}

MAX_LIST_ENTRIES = 500
MAX_SEARCH_RESULTS = 100
MAX_SEARCH_VISITS = 20_000
MAX_SEARCH_DEPTH = 8
SEARCH_SECONDS = 10.0


class AgentError(Exception):
    def __init__(self, status_code: int, code: str, message: str):
        super().__init__(message)
        self.status_code = status_code
        self.code = code
        self.message = message


def enabled() -> bool:
    return get_settings().enable_desktop_agent


def require_enabled() -> None:
    if not enabled():
        raise AgentError(403, "agent_disabled", "Desktop agent is disabled. Enable it only on your own PC "
                                               "(ATLAS_ENABLE_DESKTOP_AGENT=true).")


# --- folders --------------------------------------------------------------------

def approved_folders() -> dict[str, Path]:
    raw = (os.getenv("ATLAS_AGENT_FOLDERS") or "").strip()
    folders: dict[str, Path] = {}
    if raw:
        for item in raw.split(";"):
            name, sep, path = item.partition("=")
            name = name.strip().lower()
            if sep and name and path.strip() and name.replace("_", "").isalnum():
                folders[name] = Path(path.strip()).expanduser()
    else:
        home = Path.home()
        for name in ("Documents", "Downloads", "Desktop"):
            folders[name.lower()] = home / name
    return {name: path.resolve() for name, path in folders.items() if path.is_dir()}


def _resolve(folder: str, subpath: str = "") -> tuple[Path, Path]:
    roots = approved_folders()
    if folder not in roots:
        raise AgentError(422, "unknown_folder", f"'{folder}' is not an approved folder. Approved: {', '.join(sorted(roots)) or 'none'}.")
    root = roots[folder]
    target = (root / subpath).resolve() if subpath else root
    if target != root and root not in target.parents:
        raise AgentError(422, "outside_folder", "That path is outside the approved folder.")
    return root, target


def _rel(root: Path, path: Path) -> str:
    return path.relative_to(root).as_posix() or "."


def _iso(ts: float) -> str:
    return datetime.fromtimestamp(ts, UTC).isoformat(timespec="seconds")


# --- actions --------------------------------------------------------------------

def open_app(app_id: str) -> dict:
    require_enabled()
    app = APPS.get(app_id)
    if app is None:
        raise AgentError(422, "unknown_app", "That app is not on the allowlist.")
    system = platform.system()
    command = app["commands"].get(system)
    if not command:
        raise AgentError(501, "unsupported_os", f"{app['label']} is not configured for this operating system.")
    executable = command[0]
    if (system != "Windows" or not executable.lower().endswith(".exe")) and shutil.which(executable) is None:
        raise AgentError(503, "app_missing", f"{app['label']} is not installed or not on PATH.")
    try:
        # Fixed argument list from the allowlist; never shell=True, never user text.
        subprocess.Popen(command, shell=False, close_fds=True)  # noqa: S603
    except OSError:
        raise AgentError(503, "launch_failed", f"Could not launch {app['label']} on this machine.") from None
    return {"launched": True, "app_id": app_id, "label": app["label"]}


def list_folder(folder: str, subpath: str = "") -> dict:
    require_enabled()
    root, target = _resolve(folder, subpath)
    if not target.is_dir():
        raise AgentError(404, "not_a_folder", "That folder does not exist.")
    entries, total = [], 0
    with os.scandir(target) as it:
        for entry in it:
            total += 1
            if len(entries) >= MAX_LIST_ENTRIES:
                continue
            try:
                if entry.is_symlink():
                    kind, size, mtime = "link", None, None  # never follow links out of the folder
                else:
                    st = entry.stat(follow_symlinks=False)
                    kind = "folder" if entry.is_dir(follow_symlinks=False) else "file"
                    size, mtime = (None if kind == "folder" else st.st_size), _iso(st.st_mtime)
            except OSError:
                kind, size, mtime = "unreadable", None, None
            entries.append({"name": entry.name, "type": kind, "size": size, "modified": mtime})
    entries.sort(key=lambda e: (e["type"] != "folder", e["name"].lower()))
    return {"folder": folder, "path": _rel(root, target), "entries": entries, "total": total,
            "truncated": total > len(entries)}


def search_files(folder: str, query: str, subpath: str = "") -> dict:
    require_enabled()
    needle = query.strip().lower()
    if not needle:
        raise AgentError(422, "empty_query", "Search text must not be blank.")
    root, start = _resolve(folder, subpath)
    if not start.is_dir():
        raise AgentError(404, "not_a_folder", "That folder does not exist.")
    deadline = time.monotonic() + SEARCH_SECONDS
    results, visited, stopped = [], 0, None
    base_depth = len(start.parts)
    for dirpath, dirnames, filenames in os.walk(start, followlinks=False, onerror=lambda _e: None):
        here = Path(dirpath)
        dirs = set(dirnames)
        if len(here.parts) - base_depth >= MAX_SEARCH_DEPTH:
            dirnames[:] = []  # report these folders but do not descend further
        for name in sorted(dirs) + filenames:
            visited += 1
            if needle in name.lower():
                results.append({"path": _rel(root, here / name), "type": "folder" if name in dirs else "file"})
                if len(results) >= MAX_SEARCH_RESULTS:
                    stopped = "result limit"
                    break
        if stopped:
            break
        if visited >= MAX_SEARCH_VISITS:
            stopped = "visit limit"
            break
        if time.monotonic() > deadline:
            stopped = "time limit"
            break
    return {"folder": folder, "query": query.strip(), "results": results, "visited": visited, "stopped_at": stopped}


# --- writes: new files only, and a trash instead of delete ---------------------------

NAME_RE = re.compile(r"^[\w\- .()]{1,100}$", re.UNICODE)
MAX_NOTE_CHARS = 100_000
MEDIA_KEYS = {"play_pause": 0xB3, "next": 0xB0, "previous": 0xB1, "volume_up": 0xAF, "volume_down": 0xAE, "mute": 0xAD}


def _new_file(folder: str, subpath: str, name: str, suffix: str) -> tuple[Path, Path]:
    name = (name or "").strip()
    if not NAME_RE.match(name) or name.startswith(".") or ".." in name:
        raise AgentError(422, "bad_name", "File names may use letters, numbers, spaces, - _ . ( ) and at most 100 characters.")
    if not name.lower().endswith(suffix):
        name += suffix
    root, directory = _resolve(folder, subpath)
    if not directory.is_dir():
        raise AgentError(404, "not_a_folder", "That folder does not exist.")
    target = directory / name
    if target.exists() or target.is_symlink():
        raise AgentError(409, "file_exists", f"{_rel(root, target)} already exists. ATLAS never overwrites files; pick another name.")
    return root, target


def _write_new(target: Path, write) -> None:
    """Create the file exclusively, so a file appearing meanwhile is never overwritten."""
    tmp = target.with_name(f".atlas-{uuid.uuid4().hex[:8]}.tmp")
    try:
        write(tmp)
        os.link(tmp, target)  # fails if target exists; atomic on the same volume
    except FileExistsError:
        raise AgentError(409, "file_exists", "A file with that name appeared meanwhile; nothing was overwritten.") from None
    except OSError:
        # Some file systems (e.g. FAT/exFAT) lack hard links: fall back to an exclusive open.
        if target.exists():
            raise AgentError(409, "file_exists", "A file with that name appeared meanwhile; nothing was overwritten.") from None
        with open(target, "xb") as dest, open(tmp, "rb") as src:
            shutil.copyfileobj(src, dest)
    finally:
        tmp.unlink(missing_ok=True)


def create_note(folder: str, subpath: str, name: str, text: str) -> dict:
    require_enabled()
    if len(text) > MAX_NOTE_CHARS:
        raise AgentError(422, "too_long", f"Notes may be at most {MAX_NOTE_CHARS} characters.")
    suffix = ".md" if name.lower().endswith(".md") else ".txt"
    root, target = _new_file(folder, subpath, name, suffix)
    _write_new(target, lambda tmp: tmp.write_text(text, encoding="utf-8"))
    return {"created": _rel(root, target), "folder": folder, "bytes": target.stat().st_size}


def create_document(folder: str, subpath: str, name: str, fmt: str, title: str, text: str) -> dict:
    """A Word document from text (lines starting with # become headings) or an Excel sheet from CSV-style text."""
    require_enabled()
    if len(text) > MAX_NOTE_CHARS:
        raise AgentError(422, "too_long", f"Documents may be at most {MAX_NOTE_CHARS} characters.")
    if fmt == "docx":
        root, target = _new_file(folder, subpath, name, ".docx")

        def write(tmp: Path) -> None:
            import docx

            document = docx.Document()
            if title.strip():
                document.add_heading(title.strip()[:200], level=0)
            for line in text.splitlines():
                stripped = line.strip()
                level = len(stripped) - len(stripped.lstrip("#"))
                if 1 <= level <= 3 and stripped[level:level + 1] == " ":
                    document.add_heading(stripped[level + 1:], level=level)
                elif stripped.startswith(("- ", "* ")):
                    document.add_paragraph(stripped[2:], style="List Bullet")
                elif stripped:
                    document.add_paragraph(stripped)
            document.core_properties.comments = "Created by ATLAS OS"
            document.save(str(tmp))
    elif fmt == "xlsx":
        root, target = _new_file(folder, subpath, name, ".xlsx")

        def write(tmp: Path) -> None:
            import csv
            import io

            from openpyxl import Workbook

            book = Workbook()
            sheet = book.active
            sheet.title = (title.strip() or "Sheet1")[:31].translate(str.maketrans("", "", "[]:*?/\\"))
            for row in csv.reader(io.StringIO(text)):
                # Cells are stored as text: a value like "=HYPERLINK(...)" must never become a formula.
                sheet.append([cell for cell in row])
                for cell in sheet[sheet.max_row]:
                    cell.data_type = "s"
            book.save(str(tmp))
    else:
        raise AgentError(422, "bad_format", "Format must be docx or xlsx.")
    _write_new(target, write)
    return {"created": _rel(root, target), "folder": folder, "bytes": target.stat().st_size}


def trash_dir() -> Path:
    return storage.data_dir() / "trash"


def move_to_trash(folder: str, subpath: str) -> dict:
    require_enabled()
    if not subpath.strip():
        raise AgentError(422, "no_path", "Name the file to move to the trash.")
    root = approved_folders().get(folder)
    if root is None:
        raise AgentError(422, "unknown_folder", f"'{folder}' is not an approved folder.")
    # Check the entry itself (not where a link points): its parent must resolve inside the folder.
    entry = root / subpath
    parent = entry.parent.resolve()
    if parent != root and root not in parent.parents:
        raise AgentError(422, "outside_folder", "That path is outside the approved folder.")
    entry = parent / entry.name
    if entry.is_symlink() or not entry.is_file():
        raise AgentError(404, "not_a_file", "Only regular files can be moved to the trash (not folders or links).")
    item_id = uuid.uuid4().hex[:12]
    holder = trash_dir() / item_id
    holder.mkdir(parents=True)
    shutil.move(str(entry), str(holder / entry.name))
    record = {"id": item_id, "folder": folder, "path": _rel(root, entry), "name": entry.name,
              "trashed_at": datetime.now(UTC).isoformat(timespec="seconds")}
    (holder / "atlas-trash.json").write_text(json.dumps(record), encoding="utf-8")
    return {"trashed": record["path"], "trash_id": item_id, "restore": "Restore it from Tasks → PC actions → Restore from trash."}


def list_trash() -> list[dict]:
    items = []
    base = trash_dir()
    if base.is_dir():
        for holder in base.iterdir():
            try:
                items.append(json.loads((holder / "atlas-trash.json").read_text(encoding="utf-8")))
            except (OSError, ValueError):
                continue
    return sorted(items, key=lambda r: r.get("trashed_at", ""), reverse=True)


def restore_from_trash(trash_id: str) -> dict:
    require_enabled()
    if not re.fullmatch(r"[0-9a-f]{12}", trash_id or ""):
        raise AgentError(422, "bad_trash_id", "Unknown trash item.")
    holder = trash_dir() / trash_id
    try:
        record = json.loads((holder / "atlas-trash.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise AgentError(404, "not_in_trash", "That item is not in the trash.") from None
    root, target = _resolve(record["folder"], record["path"])
    if target.exists() or target.is_symlink():
        raise AgentError(409, "file_exists", f"{record['path']} exists again; move or rename it first.")
    if not target.parent.is_dir():
        raise AgentError(404, "not_a_folder", "The original folder no longer exists.")
    shutil.move(str(holder / record["name"]), str(target))
    shutil.rmtree(holder, ignore_errors=True)
    return {"restored": _rel(root, target), "folder": record["folder"]}


def media_key(key: str) -> dict:
    require_enabled()
    code = MEDIA_KEYS.get(key)
    if code is None:
        raise AgentError(422, "unknown_key", "Unknown media key.")
    if platform.system() != "Windows":
        raise AgentError(501, "unsupported_os", "Media keys are available on Windows only.")
    import ctypes

    user32 = ctypes.windll.user32  # type: ignore[attr-defined]
    user32.keybd_event(code, 0, 0, 0)       # key down
    user32.keybd_event(code, 0, 0x0002, 0)  # KEYEVENTF_KEYUP
    return {"pressed": key}


def index_documents(folder: str) -> dict:
    require_enabled()
    root, _ = _resolve(folder)
    return docsearch.index_folder(folder, root)


# --- registry used by the task engine and UI ----------------------------------------

ACTIONS = {
    "open_app": {"label": "Open an app", "risk": "Starts a program on this PC (allowlisted apps only)."},
    "list_folder": {"label": "List a folder", "risk": "Read-only: shows file names, sizes and dates in an approved folder."},
    "search_files": {"label": "Search file names", "risk": "Read-only: searches file and folder names in an approved folder."},
    "index_documents": {"label": "Index documents for search",
                        "risk": "Reads the text of .txt, .md, .csv, .pdf and .docx files in an approved folder into ATLAS's "
                                "local database. Nothing leaves this PC unless you turn on 'Use my documents' in chat."},
    "create_note": {"label": "Create a note", "risk": "Writes a new .txt or .md file. Never overwrites an existing file."},
    "create_document": {"label": "Create a Word or Excel file",
                        "risk": "Writes a new .docx or .xlsx file. Never overwrites an existing file."},
    "move_to_trash": {"label": "Move a file to ATLAS trash",
                      "risk": "Moves one file out of the folder into ATLAS's trash. It is not deleted and can be restored."},
    "restore_from_trash": {"label": "Restore from ATLAS trash", "risk": "Moves a trashed file back to where it was."},
    "media_key": {"label": "Media control", "risk": "Presses a media key (play/pause, next, volume). Windows only."},
}


def describe(kind: str, params: dict) -> dict:
    """Plain-language summary shown on the approval card. Built only from validated params."""
    if kind == "open_app":
        summary = f"Open {APPS.get(params.get('app_id', ''), {}).get('label', params.get('app_id'))}"
    elif kind == "list_folder":
        summary = f"List {params.get('folder')}/{params.get('path') or ''}".rstrip("/")
    elif kind == "search_files":
        summary = f"Search '{params.get('query')}' in {params.get('folder')}/{params.get('path') or ''}".rstrip("/")
    elif kind == "index_documents":
        summary = f"Index documents in {params.get('folder')}"
    elif kind in ("create_note", "create_document"):
        where = f"{params.get('folder')}/{params.get('path') or ''}".rstrip("/")
        size = len(params.get("text", ""))
        summary = f"Create '{params.get('name')}' ({params.get('format', 'note')}, {size} characters) in {where}"
    elif kind == "move_to_trash":
        summary = f"Move {params.get('folder')}/{params.get('path')} to ATLAS trash"
    elif kind == "restore_from_trash":
        summary = f"Restore trash item {params.get('trash_id')}"
    else:
        summary = f"Press media key: {params.get('key')}"
    return {"label": ACTIONS[kind]["label"], "risk": ACTIONS[kind]["risk"], "summary": summary}


def validate(kind: str, params: dict) -> None:
    """Checked when the task is created, so the user approves something that can run."""
    require_enabled()
    if kind == "open_app":
        if params.get("app_id") not in APPS:
            raise AgentError(422, "unknown_app", "That app is not on the allowlist.")
    elif kind == "media_key":
        if params.get("key") not in MEDIA_KEYS:
            raise AgentError(422, "unknown_key", "Unknown media key.")
    elif kind == "restore_from_trash":
        if not any(r["id"] == params.get("trash_id") for r in list_trash()):
            raise AgentError(404, "not_in_trash", "That item is not in the trash.")
    elif kind in ("create_note", "create_document"):
        _new_file(params.get("folder", ""), params.get("path", ""), params.get("name", ""),
                  ".docx" if params.get("format") == "docx" else ".xlsx" if params.get("format") == "xlsx"
                  else ".md" if params.get("name", "").lower().endswith(".md") else ".txt")
    else:
        _resolve(params.get("folder", ""), params.get("path", ""))


def run(kind: str, params: dict) -> dict:
    if kind == "open_app":
        return open_app(params["app_id"])
    if kind == "list_folder":
        return list_folder(params["folder"], params.get("path", ""))
    if kind == "search_files":
        return search_files(params["folder"], params["query"], params.get("path", ""))
    if kind == "index_documents":
        return index_documents(params["folder"])
    if kind == "create_note":
        return create_note(params["folder"], params.get("path", ""), params["name"], params["text"])
    if kind == "create_document":
        return create_document(params["folder"], params.get("path", ""), params["name"], params["format"],
                               params.get("title", ""), params["text"])
    if kind == "move_to_trash":
        return move_to_trash(params["folder"], params["path"])
    if kind == "restore_from_trash":
        return restore_from_trash(params["trash_id"])
    if kind == "media_key":
        return media_key(params["key"])
    raise AgentError(422, "unknown_action", "Unknown action.")


def status() -> dict:
    on = enabled()
    return {"enabled": on, "platform": platform.system(),
            "apps": [{"id": k, "label": v["label"]} for k, v in APPS.items()],
            "folders": sorted(approved_folders()) if on else [],
            "actions": [{"kind": k, **v} for k, v in ACTIONS.items()],
            "media_keys": list(MEDIA_KEYS) if on and platform.system() == "Windows" else []}
